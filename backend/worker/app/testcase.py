"""用例生成 Skill 真实执行器（M2 W8 / PRD 5.4.1）。

执行流程：
1. 读任务快照 ``case_source``（file_id 或 text）：文本类（PRD/OpenAPI/Markdown）
   经 Agent 协议档 LLM 按六策略配比生成；Excel 由 openpyxl 抽取文本行后同路生成；
2. 规模口径：目标中等 45 条，硬上限 80（PRD：简单/中/复杂 20/45/80），
   超上限 → failed 提示拆分，禁止灌水；
3. 自检红字（P0）：无核心正向、缺约束反向 → 写入 ``case_sets.checks``；
4. 落库用例集（generated）与用例行，任务转 ``awaiting_case_confirm``，
   ``expires_at = now + 72h``；确认/废弃由 api 侧 confirm/cancel 联动任务终态，
   72h 未确认由主循环扫描器取消；
5. 5 分钟总预算：LLM 调用超时 280s，超时 → failed(TIMEOUT)。

prompt 与解析口径镜像 ``api/app/routers/cases.py`` 的 ai-generate（中文策略名、
JSON 数组输出、配比校正），修改任一侧必须同步另一侧。
"""

from __future__ import annotations

import json
import logging
from datetime import datetime, timedelta, timezone
from pathlib import Path

from sqlalchemy.orm import Session

from .casegen import (
    MAX_COUNT,
    SOURCE_MAX_CHARS,
    TARGET_COUNT,
    build_prompts,
    parse_cases,
    rebalance_by_strategy,
    selfcheck,
)
from .crypto import decrypt_secret
from .db import SessionLocal
from .events import push_ws
from .models import CaseItem, CaseSet, ProtocolProfile, Setting, StoredFile, Task, TaskEvent
from .protocol import ProtocolCallError, call_protocol
from .task_state import claim_running_task_for_terminal_write

logger = logging.getLogger("worker.testcase")

# ─── 超时与确认窗口口径（PRD 5.4.1） ───
LLM_TIMEOUT_S = 280.0  # 单次 LLM 调用超时；总预算 5 分钟留 20s 落库余量
LLM_MAX_TOKENS = 8192  # 45 条用例的输出 token 预算
CONFIRM_WINDOW_H = 72  # 确认窗口：72h 未确认由扫描器取消


def _now():
    """统一 UTC 时间戳。"""
    return datetime.now(timezone.utc)


def _fail(db: Session, task: Task, code: str, message: str) -> None:
    """任务失败收尾：落终态、写错误时间线并推送 WS error 事件。"""
    task_id = task.id
    task = claim_running_task_for_terminal_write(db, task_id)
    if not task:
        logger.info("testcase task %s skipped failure because it is no longer running", task_id)
        return
    task.status = "failed"
    task.finished_at = _now()
    task.result = {**(task.result or {}), "error_code": code, "error_message": message}
    db.add(TaskEvent(task_id=task.id, event="error", level="error", message=message, payload={"code": code}))
    db.commit()
    push_ws(task.session_id, "error", {"code": code, "message": message}, task_id=task.id)
    logger.warning("testcase task %s failed (%s): %s", task.id, code, message)


def _read_source(db: Session, case_source: dict) -> str:
    """读取来源文档内容：file_id 走文件卷，text 直取；Excel 抽取文本行。

    口径镜像 api 侧 ``cases.py`` 的来源文档读取（xlsx 逐行拼接、其余 UTF-8 容错读）。
    """
    text = (case_source.get("text") or "").strip()
    if text:
        return text[:SOURCE_MAX_CHARS]
    file_id = case_source.get("file_id")
    if not file_id:
        raise ValueError("case_source 缺少 file_id 或 text")
    stored = db.query(StoredFile).filter(StoredFile.id == file_id).first()
    if not stored:
        raise FileNotFoundError("来源文件不存在或已删除")
    path = Path(stored.storage_path)
    if not path.exists():
        raise FileNotFoundError("来源文件已从磁盘移除")
    suffix = path.suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            lines = []
            for sheet in workbook.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    line = " ".join(str(cell) for cell in row if cell is not None)
                    if line.strip():
                        lines.append(line)
        finally:
            workbook.close()
        return "\n".join(lines)[:SOURCE_MAX_CHARS]
    return path.read_bytes().decode("utf-8", errors="ignore")[:SOURCE_MAX_CHARS]


def _to_case_items(case_set_id: str, cases: list[dict]) -> list[CaseItem]:
    """把模型输出的用例字典转换为待落库的 CaseItem 行（固定字段 + 透传扩展键）。"""
    fixed = {
        "strategy", "priority", "module", "name", "precondition", "steps", "expected", "test_type"
    }
    items: list[CaseItem] = []
    for index, case in enumerate(cases):
        extras = {k: v for k, v in case.items() if k not in fixed}
        items.append(
            CaseItem(
                case_set_id=case_set_id,
                strategy=str(case.get("strategy") or "正向"),
                priority=str(case.get("priority") or "P1"),
                module=str(case.get("module") or ""),
                name=str(case.get("name") or ""),
                precondition=str(case.get("precondition") or ""),
                steps=str(case.get("steps") or ""),
                expected=str(case.get("expected") or ""),
                test_type=str(case.get("test_type") or ""),
                extras=extras,
                sort_order=index,
            )
        )
    return items


def run_testcase(task_id: str) -> None:
    """testcase 任务真实执行入口：生成用例集并转入 awaiting_case_confirm。"""
    db = SessionLocal()
    task: Task | None = None
    try:
        task = db.query(Task).filter(Task.id == task_id).first()
        if not task:
            logger.error("testcase task %s not found, skip execution", task_id)
            return
        config = task.config or {}
        case_source = config.get("case_source") or {}
        if not case_source:
            _fail(db, task, "VALIDATION", "任务快照缺少 case_source，无法生成用例")
            return

        # ─── 读取来源文档 ───
        try:
            source_text = _read_source(db, case_source)
        except FileNotFoundError as exc:
            _fail(db, task, "VALIDATION", str(exc))
            return
        except ValueError as exc:
            _fail(db, task, "VALIDATION", str(exc))
            return
        except Exception as exc:
            _fail(db, task, "VALIDATION", "来源文档内容不可读，请改用文本输入")
            logger.exception("testcase task %s source read error: %s", task_id, exc)
            return
        if not source_text.strip():
            _fail(db, task, "VALIDATION", "来源文档内容为空，无法生成用例")
            return

        # ─── 解析 Agent 协议档并调用 LLM 生成（六策略配比） ───
        row = db.query(Setting).filter(Setting.key == "agent_profile_id").first()
        profile_id = row.value if row else None
        profile = (
            db.query(ProtocolProfile).filter(ProtocolProfile.id == profile_id).first()
            if profile_id
            else None
        )
        if not profile or not profile.encrypted_key:
            _fail(db, task, "VALIDATION", "未配置 Agent 协议档或未填写 API Key，无法生成用例")
            return

        db.add(TaskEvent(task_id=task.id, event="progress", message="正在按六策略生成用例"))
        db.commit()
        push_ws(task.session_id, "progress", {"percent": 20, "done": 0, "total": 1, "message": "正在生成用例"}, task_id=task.id)

        system, user = build_prompts(source_text, TARGET_COUNT)
        try:
            result = call_protocol(
                protocol=profile.protocol,
                base_url=profile.base_url,
                model=profile.model,
                api_key=decrypt_secret(profile.encrypted_key),
                messages=[{"role": "user", "content": user}],
                system=system,
                temperature=0.3,
                max_tokens=LLM_MAX_TOKENS,
                anthropic_version=profile.anthropic_version,
                timeout_s=LLM_TIMEOUT_S,
            )
        except ProtocolCallError as exc:
            _fail(db, task, exc.code, f"用例生成失败：{exc.message}")
            return
        if not (result.text or "").strip():
            _fail(db, task, "UPSTREAM", "模型返回空内容，用例生成失败")
            return

        # ─── 解析与规模闸门 ───
        try:
            cases = parse_cases(result.text)
        except (ValueError, json.JSONDecodeError) as exc:
            _fail(db, task, "UPSTREAM", f"模型输出无法解析为用例：{exc}")
            return
        if len(cases) > MAX_COUNT:
            _fail(
                db, task, "VALIDATION",
                f"生成 {len(cases)} 条超过上限 {MAX_COUNT} 条，请拆分需求文档后重新发起",
            )
            return
        cases = rebalance_by_strategy(cases, min(len(cases), TARGET_COUNT))
        if not cases:
            _fail(db, task, "UPSTREAM", "模型输出经配比校正后无有效用例")
            return

        # ─── 落库用例集 + 自检红字 + 任务转 awaiting_case_confirm ───
        # 模型调用期间可能收到取消；落库前重新锁定任务，不能覆盖 cancelled 终态。
        task = claim_running_task_for_terminal_write(db, task.id)
        if not task:
            logger.info("testcase task %s skipped persistence because it is no longer running", task_id)
            return
        expires_at = _now() + timedelta(hours=CONFIRM_WINDOW_H)
        case_set = CaseSet(
            task_id=task.id,
            name=f"AI 生成用例集 · 任务 {task.id[:8]}",
            status="generated",
            generated_count=len(cases),
            confirmed_count=0,
            checks=selfcheck(cases),
            expires_at=expires_at,
        )
        db.add(case_set)
        db.flush()
        db.add_all(_to_case_items(case_set.id, cases))

        task.status = "awaiting_case_confirm"
        task.progress = {"percent": 100, "done": len(cases), "total": len(cases), "message": "用例已生成，等待确认入库"}
        task.result = {"case_set_id": case_set.id, "generated_count": len(cases)}
        db.add(
            TaskEvent(
                task_id=task.id,
                event="awaiting_case_confirm",
                message=f"已生成 {len(cases)} 条用例，等待确认入库（72h 内有效）",
                payload={"case_set_id": case_set.id},
            )
        )
        db.commit()

        push_ws(
            task.session_id,
            "progress",
            {"percent": 100, "done": len(cases), "total": len(cases), "message": "用例已生成，等待确认入库"},
            task_id=task.id,
        )
        push_ws(
            task.session_id,
            "thought",
            {"text": f"已生成 {len(cases)} 条用例（含自检结果），请到「用例」页确认入库或拒绝；72 小时内未确认将自动取消。"},
            task_id=task.id,
        )
        logger.info("testcase task %s awaiting confirm (case_set=%s, %d cases)", task_id, case_set.id, len(cases))
    except Exception as exc:
        db.rollback()
        logger.exception("testcase task %s execution error", task_id)
        if task is not None:
            _fail(db, task, "INTERNAL", f"执行失败({type(exc).__name__})，详情见服务端日志")
    finally:
        db.close()
