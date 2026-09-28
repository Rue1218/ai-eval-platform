"""用例工作台缺陷回归：真实 ORM 覆盖来源权限、终态、映射及保存契约。"""

import asyncio
import json
from datetime import timedelta
from io import BytesIO
from types import SimpleNamespace

import pytest
from fastapi import UploadFile
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from shared.case_design import parse_design
from sqlalchemy import JSON, BigInteger, Integer, MetaData, create_engine, event, update
from sqlalchemy.dialects.postgresql import JSONB
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.db import get_db
from app.deps import get_current_user
from app.errors import AppError, ErrorCode
from app.harness.execution.task_tools import cancel_task_safe
from app.models import (
    AuditLog,
    CaseFolder,
    CaseItem,
    CaseSet,
    Dataset,
    DatasetRow,
    StoredFile,
    Task,
    TaskEvent,
    User,
    utcnow,
)
from app.routers import cases as routes
from app.routers import tasks as task_routes
from app.schemas import (
    CaseAiDesignIn,
    CaseAiFillIn,
    CaseAiGenerateIn,
    CaseCancelIn,
    CaseConfirmIn,
    CaseMapIn,
    CaseSetFromCandidatesIn,
    CaseSetUpdate,
    CasesPayload,
)


@pytest.fixture
def cases_db():
    """仅在内存库复制相关 DDL；保留 ORM 行为与事务，适配 JSONB 和自增主键。"""
    metadata = MetaData()
    pending = [model.__table__ for model in (
        User, CaseFolder, CaseSet, CaseItem, Task, TaskEvent, Dataset, DatasetRow, StoredFile, AuditLog,
    )]
    while pending:
        table = pending.pop()
        if table.name in metadata.tables:
            continue
        copied = table.to_metadata(metadata)
        for column in copied.columns:
            if isinstance(column.type, JSONB):
                column.type = JSON()
            elif column.primary_key and isinstance(column.type, BigInteger):
                column.type = Integer()
        for index in copied.indexes:
            condition = index.dialect_options["postgresql"].get("where")
            if condition is not None:
                index.dialect_options["sqlite"]["where"] = condition
        pending.extend(foreign.column.table for foreign in table.foreign_keys)
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    metadata.create_all(engine)
    try:
        with sessionmaker(engine, autoflush=False)() as db:
            db.add_all([
                User(id="owner", username="owner", password_hash="fake"),
                User(id="other", username="other", password_hash="fake"),
            ])
            db.add(CaseSet(id="set", name="需求用例", created_by="owner"))
            db.commit()
            yield db
    finally:
        engine.dispose()


def _actor():
    """返回固定测试成员，避免读取真实登录状态。"""
    return SimpleNamespace(id="owner")


def _request():
    """审计测试不包含真实客户端地址。"""
    return SimpleNamespace(client=None)


def _add_case(db, case_id="case", **overrides):
    """插入真实用例，缺省为完整的核心正向样本。"""
    row = CaseItem(**{
        "id": case_id, "case_set_id": "set", "name": case_id,
        "strategy": "正向", "priority": "HX", "expected": "成功", **overrides,
    })
    db.add(row)
    db.commit()
    return row


def _waiting_task(db, *, expired=False):
    """为用例集关联待确认任务，期限可由测试切换。"""
    task = Task(id="task", kind="testcase", created_by="owner", status="awaiting_case_confirm")
    db.add(task)
    case_set = db.get(CaseSet, "set")
    case_set.task_id = task.id
    case_set.expires_at = utcnow() + timedelta(hours=-1 if expired else 1)
    db.commit()
    return task


def test_generate_rejects_foreign_file_before_model_call(cases_db, monkeypatch):
    """知道私有附件 ID 不能代替来源授权，拒绝时不读取磁盘也不调用模型。"""
    cases_db.add(StoredFile(
        id="private", filename="private.txt", uploaded_by="other", size_bytes=1,
        sha256="0" * 64, storage_path="/must-not-read/private.bin",
    ))
    cases_db.commit()
    monkeypatch.setattr(routes, "call_agent_model", lambda *a, **k: pytest.fail("不能调用模型"))
    with pytest.raises(AppError) as error:
        routes.ai_generate_cases(CaseAiGenerateIn(source_doc_id="private"), cases_db, _actor())
    assert error.value.code == ErrorCode.UNAUTHORIZED
    with pytest.raises(AppError) as design_error:
        routes.ai_design_cases(CaseAiDesignIn(source_doc_id="private"), cases_db, _actor())
    assert design_error.value.code == ErrorCode.UNAUTHORIZED


def _design_json():
    """分析样本携带真实需求引文，刻意不提供平台生成的稳定 ID。"""
    return json.dumps({"summary": "登录", "test_points": [{"title": "登录成功",
                      "source_quote": "登录需求", "risk": "high", "strategies": ["positive"]}]})


def test_design_only_returns_traceable_preview_without_writing(cases_db, monkeypatch):
    """独立设计只读来源并调用分析技能，不创建空用例集或后台任务。"""
    calls = []

    def model(_db, system, _prompt, **_kwargs):
        calls.append(system)
        return SimpleNamespace(text=_design_json())

    monkeypatch.setattr(routes, "call_agent_model", model)
    result = routes.ai_design_cases(CaseAiDesignIn(source_text="登录需求"), cases_db, _actor())
    assert result["design"]["test_points"][0]["id"] == "TP-001"
    assert result["loaded_sections"] == ["workflow", "design"]
    assert "人工用例结构" not in calls[0]
    assert cases_db.query(CaseSet).count() == 1 and cases_db.query(Task).count() == 0


def test_generation_rejects_changed_source_before_model_and_filters_unselected_points(cases_db, monkeypatch):
    """需求变更先拒绝调用；通过校验后只采纳有真实测试点关联的去重候选。"""
    design = parse_design(_design_json(), "登录需求")
    monkeypatch.setattr(routes, "call_agent_model", lambda *a, **k: pytest.fail("来源变化不能调用模型"))
    with pytest.raises(AppError) as error:
        routes.ai_generate_cases(CaseAiGenerateIn(source_text="修改需求", design=design), cases_db, _actor())
    assert error.value.code == ErrorCode.VALIDATION
    case = {"name": "登录", "strategy": "正向", "test_point_id": "TP-001", "requirement_quote": "伪造"}
    output = [case, case, {**case, "test_point_id": "TP-002"}]
    monkeypatch.setattr(routes, "call_agent_model", lambda *a, **k: SimpleNamespace(text=json.dumps(output)))
    result = routes.ai_generate_cases(CaseAiGenerateIn(source_text="登录需求", design=design), cases_db, _actor())
    assert result["items"] == [{**case, "requirement_quote": "登录需求", "risk": "high"}]


def test_trace_columns_export_from_old_draft_without_registered_columns(cases_db):
    """追加到旧草稿后，原列配置为空也不能使 Excel 丢失需求依据。"""
    item = _add_case(cases_db, extras={"test_point_id": "TP-001", "requirement_quote": "登录需求", "risk": "high"})
    data = routes._build_xlsx(cases_db.get(CaseSet, "set"), [item])
    workbook = load_workbook(BytesIO(data), read_only=True)
    try:
        headers, values = list(workbook.active.values)
        exported = dict(zip(headers, values))
        assert exported["测试点编号"] == "TP-001"
        assert exported["需求原文依据"] == "登录需求"
    finally:
        workbook.close()


def test_skill_discovery_http_routes_and_source_validation():
    """固定路由不被用例集 ID 吞掉，未知技能与异常设计使用标准校验错误。"""
    from app.main import app

    app.dependency_overrides[get_current_user] = _actor
    try:
        client = TestClient(app)
        result = client.get("/api/case-sets/generation-skills")
        assert result.status_code == 200
        assert result.json()["items"][0]["id"] == "functional-test-design"
        invalid = client.post("/api/case-sets/ai-design", json={"source_text": "需求", "skill_id": "../../.env"})
        assert invalid.status_code == 400
    finally:
        app.dependency_overrides.pop(get_current_user, None)


def test_export_keeps_formula_like_requirements_and_headers_as_text(cases_db):
    """模型、需求原文与自定义列标题均是文本，不能被 Excel 解释成公式。"""
    case_set = cases_db.get(CaseSet, "set")
    case_set.column_schema = [{"key": "custom", "name": "=1+1"}]
    item = _add_case(cases_db, name="=2+2", extras={"requirement_quote": "=3+3", "custom": "=4+4"})
    workbook = load_workbook(BytesIO(routes._build_xlsx(case_set, [item])))
    try:
        literal_cells = [cell for row in workbook.active for cell in row
                         if isinstance(cell.value, str) and cell.value.startswith("=")]
        assert {cell.value for cell in literal_cells} == {"=1+1", "=2+2", "=3+3", "=4+4"}
        assert all(cell.data_type == "s" for cell in literal_cells)
    finally:
        workbook.close()


@pytest.mark.parametrize("endpoint,metadata", [
    ("generation-skills", None), ("generation-skills", "---\nname: broken\n---\n"),
    ("ai-design", "---\ndescription: test\nlicense: MIT\n---\nworkflow"),
    ("ai-generate", "---\ndescription: test\nlicense: MIT\n---\nworkflow"),
])
def test_missing_skill_resources_return_actionable_safe_error(cases_db, monkeypatch, tmp_path, endpoint, metadata):
    """资源损坏在模型调用前报明确错误，不能泄漏服务器路径或调用模型浪费额度。"""
    from shared import case_skill

    from app.main import app

    monkeypatch.setattr(case_skill, "SKILL_ROOT", tmp_path)
    if metadata is not None:
        (tmp_path / "SKILL.md").write_text(metadata, encoding="utf-8")
    monkeypatch.setattr(routes, "call_agent_model", lambda *a, **k: pytest.fail("技能不可用不能调用模型"))
    app.dependency_overrides[get_current_user] = _actor
    app.dependency_overrides[get_db] = lambda: cases_db
    try:
        client = TestClient(app, raise_server_exceptions=False)
        result = (client.get(f"/api/case-sets/{endpoint}") if endpoint == "generation-skills"
                  else client.post(f"/api/case-sets/{endpoint}", json={"source_text": "登录需求"}))
        assert result.status_code == 500
        assert result.json() == {"code": "INTERNAL", "message": "用例生成技能不可用，请联系管理员检查技能资源"}
    finally:
        app.dependency_overrides.pop(get_current_user, None)
        app.dependency_overrides.pop(get_db, None)


def test_excel_source_uses_original_filename_with_bin_storage(cases_db, tmp_path):
    """模拟上传器的 UUID.bin 存储，抽取的素材必须是单元格正文。"""
    workbook = Workbook()
    workbook.active.append(["登录需求", "错误密码应拒绝登录"])
    buffer = BytesIO()
    workbook.save(buffer)
    path = tmp_path / "upload.bin"
    path.write_bytes(buffer.getvalue())
    cases_db.add(StoredFile(
        id="excel", filename="需求.xlsx", kind="xlsx", uploaded_by="owner",
        size_bytes=path.stat().st_size, sha256="0" * 64, storage_path=str(path),
    ))
    cases_db.commit()
    assert routes._read_stored_file_text(cases_db, "excel", "owner") == "登录需求 错误密码应拒绝登录"


def test_generate_applies_requested_weights_and_drops_model_identity(cases_db, monkeypatch):
    """全正向配置贯穿提示词和后处理，供应商不能指定待保存行的身份。"""
    output = [{"name": str(index), "strategy": "正向", "id": "forged"} for index in range(20)]
    output.append({"name": "多余反向", "strategy": "反向"})
    monkeypatch.setattr(routes, "call_agent_model", lambda *a, **k: SimpleNamespace(text=json.dumps(output)))
    result = routes.ai_generate_cases(
        CaseAiGenerateIn(source_text="需求", max_count=20, strategy_weights={"positive": 100}),
        cases_db, _actor(),
    )
    assert len(result["items"]) == 20
    assert all("id" not in item and item["strategy"] == "正向" for item in result["items"])


@pytest.mark.parametrize("max_count,budget", [(45, 8192), (80, 16384)])
def test_generate_uses_shared_output_budget(cases_db, monkeypatch, max_count, budget):
    """页面候选生成按目标条数传递共享输出 token 预算。"""
    calls = []

    def fake_model(*_args, **kwargs):
        calls.append(kwargs)
        return SimpleNamespace(text='[{"name":"登录","strategy":"正向","priority":"HX"}]')

    monkeypatch.setattr(routes, "call_agent_model", fake_model)
    routes.ai_generate_cases(CaseAiGenerateIn(source_text="登录需求", max_count=max_count), cases_db, _actor())
    assert calls[0]["max_tokens"] == budget


@pytest.mark.parametrize("text,raw,expected", [
    ("", {"choices": [{"message": {"reasoning_content": "private"}}]}, "仅返回思考"),
    ('[{"name":"登录","strategy":"正向"}]', {"stop_reason": "max_tokens"}, "长度上限"),
])
def test_generate_rejects_unfinished_body(cases_db, monkeypatch, text, raw, expected):
    """页面生成与 Worker 共用失败口径，截断的可解析片段也不能冒充完整结果。"""
    monkeypatch.setattr(routes, "call_agent_model", lambda *a, **k: SimpleNamespace(text=text, raw=raw))
    with pytest.raises(AppError) as error:
        routes.ai_generate_cases(CaseAiGenerateIn(source_text="登录需求"), cases_db, _actor())
    assert error.value.code == ErrorCode.UPSTREAM
    assert expected in error.value.message and "private" not in error.value.message


def test_generate_rejects_when_selected_strategy_has_no_cases(cases_db, monkeypatch):
    """模型只返回未选或未知策略时，不能把空候选伪装成成功。"""
    output = [{"name": "反向", "strategy": "反向"}, {"name": "未知", "strategy": "其他"}]
    monkeypatch.setattr(routes, "call_agent_model", lambda *a, **k: SimpleNamespace(text=json.dumps(output)))
    with pytest.raises(AppError) as error:
        routes.ai_generate_cases(
            CaseAiGenerateIn(source_text="需求", strategy_weights={"positive": 100}),
            cases_db, _actor(),
        )
    assert error.value.code == ErrorCode.UPSTREAM


def test_candidates_create_one_complete_draft(cases_db):
    """采纳候选后同一事务写入草稿、用例和审计，身份字段由服务端生成。"""
    result = routes.create_case_set_from_candidates(CaseSetFromCandidatesIn(
        name=" 登录用例 ", cases=[
            {"id": "model-forged", "name": "登录成功", "strategy": "正向", "priority": "HX",
             "module": "登录", "feature_point": "正确凭据", "steps": "输入有效凭据", "expected": "进入首页"},
            {"name": "登录失败", "strategy": "反向", "priority": "YC",
             "submodule": "密码", "test_type": "异常处理", "expected": "拒绝登录"},
        ],
    ), _request(), cases_db, _actor())
    created = cases_db.get(CaseSet, result.id)
    assert created.name == "登录用例" and created.status == "generated"
    assert created.revision == 1 and created.generated_count == 2 and created.checks == []
    assert [case["name"] for case in result.cases] == ["登录成功", "登录失败"]
    assert result.cases[0]["id"] != "model-forged"
    assert result.cases[0]["feature_point"] == "正确凭据"
    assert result.cases[1]["test_type"] == "异常处理"
    assert cases_db.query(AuditLog).filter_by(target_id=result.id).count() == 1


def test_candidates_http_create_preserves_extension_in_export(cases_db):
    """真实 HTTP 成功响应可序列化，候选扩展列随草稿进入详情和 Excel。"""
    from app.main import app

    app.dependency_overrides[get_db] = lambda: cases_db
    app.dependency_overrides[get_current_user] = _actor
    try:
        client = TestClient(app)
        response = client.post("/api/case-sets/from-candidates", json={
            "name": "需求用例", "cases": [
                {"name": "登录成功", "strategy": "正向", "priority": "HX",
                 "expected": "进入首页", "requirement_id": "REQ-1", "mapped": True},
            ],
        })
        assert response.status_code == 201, response.text
        created = response.json()
        assert created["revision"] == 1 and created["generated_count"] == 1
        assert created["cases"][0]["requirement_id"] == "REQ-1"
        assert created["cases"][0]["mapped"] is False
        assert [column["key"] for column in created["column_schema"]] == ["requirement_id"]

        detail = client.get(f"/api/case-sets/{created['id']}")
        assert detail.status_code == 200
        assert detail.json()["cases"][0]["requirement_id"] == "REQ-1"
        exported = client.get(f"/api/case-sets/{created['id']}/export?fmt=xlsx")
        assert exported.status_code == 200
        workbook = load_workbook(BytesIO(exported.content), read_only=True, data_only=True)
        try:
            rows = list(workbook.active.values)
        finally:
            workbook.close()
        assert rows[0][-1] == "requirement_id"
        assert rows[1][-1] == "REQ-1"
    finally:
        app.dependency_overrides.pop(get_db, None)
        app.dependency_overrides.pop(get_current_user, None)


def test_candidates_write_failure_rolls_back_entire_draft(cases_db, monkeypatch):
    """用例行写入失败时不留下空集或半写入审计。"""
    write = routes._upsert_case

    def fail_after_first(*args, **kwargs):
        write(*args, **kwargs)
        raise RuntimeError("模拟写入中断")

    monkeypatch.setattr(routes, "_upsert_case", fail_after_first)
    with pytest.raises(RuntimeError):
        routes.create_case_set_from_candidates(CaseSetFromCandidatesIn(
            name="失败候选集", cases=[{"name": "候选", "strategy": "正向", "priority": "HX"}],
        ), _request(), cases_db, _actor())
    assert cases_db.query(CaseSet).count() == 1
    assert cases_db.query(CaseItem).count() == 0
    assert cases_db.query(AuditLog).count() == 0


def test_save_reorders_ids_recomputes_checks_and_ignores_workflow_fields(cases_db):
    """保存响应跟随提交顺序，真实状态不受客户端平铺的保留字段覆盖。"""
    _add_case(cases_db, "first", sort_order=0)
    _add_case(cases_db, "second", sort_order=1, strategy="反向")
    result = routes.save_cases("set", CasesPayload(expected_revision=0, cases=[
        {"id": "second", "name": "反向", "strategy": "反向", "priority": "YC", "expected": "拒绝"},
        {"id": "first", "name": "正向", "strategy": "正向", "priority": "HX", "expected": "成功", "mapped": True},
        {"name": "新增", "strategy": "边界", "priority": "BJ", "expected": "边界有效"},
    ]), cases_db, _actor())
    assert [item["id"] for item in result["items"][:2]] == ["second", "first"]
    assert result["items"][2]["id"] not in {"first", "second"}
    assert result["items"][1]["mapped"] is False
    assert result["checks"] == []
    assert cases_db.get(CaseSet, "set").generated_count == 3
    assert result["revision"] == 1


def test_save_empty_clears_last_row_and_recomputes_checks(cases_db):
    """最后一条用例删除后允许提交空快照，计数和检查项跟随真实剩余行。"""
    _add_case(cases_db)
    result = routes.save_cases("set", CasesPayload(expected_revision=0, cases=[]), cases_db, _actor())
    assert result["items"] == [] and result["total"] == 0
    assert cases_db.get(CaseSet, "set").generated_count == 0
    assert len(result["checks"]) == 2


def test_stale_snapshot_cannot_delete_another_save(cases_db):
    """两人读取同一修订号后，旧快照不能删除先保存的新用例。"""
    _add_case(cases_db, "first")
    original = routes.get_case_set("set", cases_db, _actor())
    assert original.revision == 0
    newer = routes.save_cases("set", CasesPayload(expected_revision=original.revision, cases=[
        {"id": "first", "name": "first", "strategy": "正向", "priority": "HX", "expected": "成功"},
        {"name": "new", "strategy": "边界", "priority": "BJ", "expected": "有效"},
    ]), cases_db, _actor())
    with pytest.raises(AppError) as error:
        routes.save_cases("set", CasesPayload(expected_revision=original.revision, cases=[
            {"id": "first", "name": "旧页面", "strategy": "正向", "priority": "HX", "expected": "成功"},
        ]), cases_db, _actor())
    assert error.value.code == ErrorCode.CONCURRENCY
    assert newer["revision"] == cases_db.get(CaseSet, "set").revision == 1
    assert cases_db.query(CaseItem).count() == 2
    assert cases_db.get(CaseItem, "first").name == "first"


def test_empty_draft_cannot_be_confirmed(cases_db):
    """空用例集可暂存草稿，但审核入库至少需要一条用例。"""
    with pytest.raises(AppError) as error:
        routes.confirm_case_set("set", CaseConfirmIn(ok=True, expected_revision=0),
                                _request(), cases_db, _actor())
    assert error.value.code == ErrorCode.VALIDATION
    assert cases_db.get(CaseSet, "set").status == "generated"


def test_save_columns_and_rows_conflict_without_partial_metadata(cases_db):
    """旧修订号的列和用例行必须一起拒绝，不出现半次保存。"""
    newer = routes.save_cases("set", CasesPayload(expected_revision=0, cases=[], column_schema=[
        {"key": "owner", "name": "负责人", "type": "string"},
    ]), cases_db, _actor())
    assert newer["revision"] == 1
    with pytest.raises(AppError) as error:
        routes.save_cases("set", CasesPayload(expected_revision=0, cases=[], column_schema=[
            {"key": "team", "name": "团队", "type": "string"},
        ]), cases_db, _actor())
    assert error.value.code == ErrorCode.CONCURRENCY
    assert cases_db.get(CaseSet, "set").column_schema[0]["key"] == "owner"


def test_independent_column_update_invalidates_stale_cases(cases_db):
    """旧元信息接口修改列后也应让旧行快照保存收到冲突。"""
    updated = routes.update_case_set("set", CaseSetUpdate(expected_revision=0, column_schema=[
        {"key": "owner", "name": "负责人", "type": "string"},
    ]), _request(), cases_db, _actor())
    assert updated.revision == 1
    with pytest.raises(AppError) as error:
        routes.save_cases("set", CasesPayload(expected_revision=0, cases=[]), cases_db, _actor())
    assert error.value.code == ErrorCode.CONCURRENCY


def test_stale_independent_column_update_does_not_overwrite(cases_db):
    """旧页面独立提交列定义时同样必须校验修订号。"""
    cases_db.add(CaseFolder(id="folder", name="原目录"))
    cases_db.get(CaseSet, "set").folder_id = "folder"
    cases_db.commit()
    routes.save_cases("set", CasesPayload(expected_revision=0, cases=[], column_schema=[
        {"key": "owner", "name": "负责人", "type": "string"},
    ]), cases_db, _actor())
    with pytest.raises(AppError) as error:
        routes.update_case_set("set", CaseSetUpdate(expected_revision=0, folder_id=None, column_schema=[
            {"key": "team", "name": "团队", "type": "string"},
        ]), _request(), cases_db, _actor())
    assert error.value.code == ErrorCode.CONCURRENCY
    assert cases_db.get(CaseSet, "set").column_schema[0]["key"] == "owner"
    assert cases_db.get(CaseSet, "set").folder_id == "folder"


def test_independent_column_update_rejects_null_without_db_error(cases_db):
    """显式 null 应返回业务校验错误，不能写入非空列或推进修订号。"""
    with pytest.raises(AppError) as error:
        routes.update_case_set("set", CaseSetUpdate(expected_revision=0, column_schema=None),
                               _request(), cases_db, _actor())
    assert error.value.code == ErrorCode.VALIDATION
    assert cases_db.get(CaseSet, "set").column_schema == []
    assert cases_db.get(CaseSet, "set").revision == 0


@pytest.mark.parametrize("status", ["confirmed", "cancelled"])
def test_terminal_set_rejects_save_and_ai_fill(cases_db, monkeypatch, status):
    """终态同时约束持久编辑和模型补全，不能继续消耗供应商调用。"""
    _add_case(cases_db)
    cases_db.get(CaseSet, "set").status = status
    cases_db.commit()
    monkeypatch.setattr(routes, "call_agent_model", lambda *a, **k: pytest.fail("终态不能补全"))
    with pytest.raises(AppError):
        routes.save_cases("set", CasesPayload(expected_revision=0, cases=[]), cases_db, _actor())
    with pytest.raises(AppError):
        routes.ai_fill_cases("set", CaseAiFillIn(case_ids=["case"]), cases_db, _actor())
    assert cases_db.query(CaseItem).count() == 1


def test_confirm_expired_set_cannot_succeed_before_expiry_scan(cases_db):
    """扫描尚未执行也不能确认过期结果，任务和用例集不会出现相反终态。"""
    task = _waiting_task(cases_db, expired=True)
    with pytest.raises(AppError) as error:
        routes.confirm_case_set("set", CaseConfirmIn(ok=True, expected_revision=0), _request(), cases_db, _actor())
    assert error.value.code == ErrorCode.VALIDATION
    assert cases_db.get(CaseSet, "set").status == "generated"
    assert task.status == "awaiting_case_confirm"


def test_confirm_rejects_revision_not_reviewed_by_user(cases_db):
    """审核页读取后新增用例时，不可直接确认用户未看过的最新版本。"""
    task = _waiting_task(cases_db)
    routes.save_cases("set", CasesPayload(expected_revision=0, cases=[
        {"name": "新用例", "strategy": "正向", "priority": "HX", "expected": "成功"},
    ]), cases_db, _actor())
    with pytest.raises(AppError) as error:
        routes.confirm_case_set("set", CaseConfirmIn(ok=True, expected_revision=0),
                                _request(), cases_db, _actor())
    assert error.value.code == ErrorCode.CONCURRENCY
    assert task.status == "awaiting_case_confirm"
    assert cases_db.get(CaseSet, "set").status == "generated"


def test_confirm_uses_case_set_then_task_locks_and_refreshes_stale_state(cases_db):
    """记录真实 ORM 的加锁语句，并证明旧 identity map 不会覆盖已取消的任务。"""
    task = _waiting_task(cases_db)
    _ = task.status
    cases_db.execute(update(Task).where(Task.id == task.id).values(status="cancelled"),
                     execution_options={"synchronize_session": False})
    locked = []

    def record_lock(execute_state):
        """SQLite 不模拟 PostgreSQL 行锁，仅检查生产查询锁序及刷新行为。"""
        statement = execute_state.statement
        if execute_state.is_select and statement._for_update_arg is not None:
            locked.append(statement.column_descriptions[0]["entity"])

    event.listen(cases_db, "do_orm_execute", record_lock)
    with pytest.raises(AppError):
        routes.confirm_case_set("set", CaseConfirmIn(ok=True, expected_revision=0), _request(), cases_db, _actor())
    assert locked == [CaseSet, Task]
    assert task.status == "cancelled"
    assert cases_db.get(CaseSet, "set").status == "generated"


def test_cancel_transitions_case_set_and_task_together(cases_db):
    """废弃用例与关联任务在同一事务产生一致终态。"""
    task = _waiting_task(cases_db)
    result = routes.cancel_case_set("set", CaseCancelIn(expected_revision=0), _request(), cases_db, _actor())
    assert result.status == task.status == "cancelled"
    assert task.finished_at is not None


def test_stale_cancel_cannot_discard_unreviewed_revision(cases_db):
    """同事保存新版后，旧页面的废弃请求不得废弃新版草稿。"""
    task = _waiting_task(cases_db)
    routes.save_cases("set", CasesPayload(expected_revision=0, cases=[
        {"name": "新用例", "strategy": "正向", "priority": "HX", "expected": "成功"},
    ]), cases_db, _actor())
    with pytest.raises(AppError) as error:
        routes.cancel_case_set("set", CaseCancelIn(expected_revision=0), _request(), cases_db, _actor())
    assert error.value.code == ErrorCode.CONCURRENCY
    assert task.status == "awaiting_case_confirm"
    assert cases_db.get(CaseSet, "set").status == "generated"


def test_task_cancel_transitions_draft_with_case_set_first_lock(cases_db):
    """任务中心取消待审核任务时同步废弃草稿，锁序与确认流程相同。"""
    task = _waiting_task(cases_db)
    locked = []

    def record_lock(execute_state):
        """SQLite 不执行行锁，用 SQLAlchemy 查询标记记录生产锁序。"""
        statement = execute_state.statement
        if execute_state.is_select and statement._for_update_arg is not None:
            locked.append(statement.column_descriptions[0]["entity"])

    event.listen(cases_db, "do_orm_execute", record_lock)
    result = task_routes.cancel_task(task.id, _request(), cases_db, _actor())
    assert result["status"] == "cancelled"
    assert cases_db.get(CaseSet, "set").status == "cancelled"
    assert locked[:2] == [CaseSet, Task]


def test_running_stress_cancel_records_external_stop(cases_db):
    """旧运行中压测取消后保留待停发标记，供 Worker 重试外部引擎。"""
    cases_db.add(Task(id="stress", kind="stress", created_by="owner", status="running", result={}))
    cases_db.commit()
    result = task_routes.cancel_task("stress", _request(), cases_db, _actor())
    assert result["status"] == "cancelled"
    assert cases_db.get(Task, "stress").result["stress_stop_pending"] is True


def test_agent_task_cancel_also_discards_generated_draft(cases_db, monkeypatch):
    """Agent 工具入口与任务中心使用同一个取消状态联动。"""
    from app import db as db_module

    _waiting_task(cases_db)
    monkeypatch.setattr(db_module, "SessionLocal", lambda: cases_db)
    result = cancel_task_safe({"task_id": "task"}, SimpleNamespace(session_id="session", user_id="owner"))
    assert result["status"] == "cancelled"
    assert cases_db.get(CaseSet, "set").status == "cancelled"


def test_map_rejects_published_dataset_without_legacy_write(cases_db):
    """评测映射已经停用，任何目标数据集都不能再写入。"""
    _add_case(cases_db)
    cases_db.add(Dataset(id="dataset", name="已发布", active_version_id="version"))
    cases_db.commit()
    with pytest.raises(AppError) as error:
        routes.map_cases("set", CaseMapIn(target="dataset", target_id="dataset", case_ids=["case"]),
                         _request(), cases_db, _actor())
    assert "已停用" in error.value.message
    assert cases_db.query(DatasetRow).count() == 0
    assert cases_db.get(CaseItem, "case").mapped is False


def test_map_rejects_generated_draft_without_legacy_write(cases_db):
    """未审核草稿不可通过旧映射接口进入评测数据集。"""
    _add_case(cases_db)
    cases_db.add(Dataset(id="dataset", name="旧数据集"))
    cases_db.commit()
    body = CaseMapIn(target="dataset", target_id="dataset", case_ids=["case", "case"])
    with pytest.raises(AppError) as error:
        routes.map_cases("set", body, _request(), cases_db, _actor())
    assert error.value.code == ErrorCode.VALIDATION
    assert cases_db.query(DatasetRow).count() == 0
    assert cases_db.get(CaseItem, "case").mapped is False


def test_ai_fill_rejects_foreign_ids_and_filters_metadata(cases_db, monkeypatch):
    """模型不能改写已有名称、非请求字段或流程状态，只能补当前行缺失内容。"""
    _add_case(cases_db, expected="")
    output = [
        {"id": "foreign", "expected": "越权"},
        {"id": "case", "name": "改写原名", "expected": "补全", "mapped": True, "case_set_id": "foreign"},
    ]
    monkeypatch.setattr(routes, "call_agent_model", lambda *a, **k: SimpleNamespace(text=json.dumps(output)))
    result = routes.ai_fill_cases("set", CaseAiFillIn(case_ids=["case"]), cases_db, _actor())
    assert result == {"items": [{"id": "case", "expected": "补全"}]}
    assert cases_db.get(CaseItem, "case").expected == ""


def test_ai_fill_rechecks_state_after_model_returns(cases_db, monkeypatch):
    """生成候选期间发生废弃，返回前再次拒绝，行锁不会覆盖整个供应商调用。"""
    _add_case(cases_db, expected="")

    def model_returns_after_cancel(*args, **kwargs):
        """模拟外部终态已提交，不调用真实供应商。"""
        cases_db.get(CaseSet, "set").status = "cancelled"
        cases_db.commit()
        return SimpleNamespace(text='[{"id":"case","expected":"不能继续编辑"}]')

    monkeypatch.setattr(routes, "call_agent_model", model_returns_after_cancel)
    with pytest.raises(AppError) as error:
        routes.ai_fill_cases("set", CaseAiFillIn(case_ids=["case"]), cases_db, _actor())
    assert error.value.code == ErrorCode.VALIDATION
    assert cases_db.get(CaseItem, "case").expected == ""


def test_append_import_recomputes_checks_from_entire_set(cases_db):
    """追加反向用例后检查项同时包含已有核心正向，不能仅检查本次导入批次。"""
    _add_case(cases_db)
    workbook = Workbook()
    workbook.active.append(["策略", "优先级", "模块", "用例名称", "预期结果"])
    workbook.active.append(["反向", "YC", "登录", "错误密码", "拒绝登录"])
    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    result = asyncio.run(routes.import_case_set(
        "set", _request(), UploadFile(file=buffer, filename="补充.xlsx"), "append", cases_db, _actor(),
        expected_revision=0,
    ))
    assert result["generated_count"] == 2
    assert result["checks"] == []
    assert result["revision"] == 1
    with pytest.raises(AppError) as error:
        routes.save_cases("set", CasesPayload(expected_revision=0, cases=[]), cases_db, _actor())
    assert error.value.code == ErrorCode.CONCURRENCY
    assert cases_db.query(CaseItem).count() == 2


def test_create_and_import_excel_commits_one_complete_set(cases_db):
    """新集导入只产生含用例的一个集，并将创建和导入审计同事务入库。"""
    workbook = Workbook()
    workbook.active.append(["用例名称", "预期结果"])
    workbook.active.append(["登录成功", "进入首页"])
    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)

    result = asyncio.run(routes.create_case_set_from_excel(
        _request(), UploadFile(file=buffer, filename="需求.xlsx"), "新导入集", None, cases_db, _actor(),
    ))

    created = cases_db.get(CaseSet, result["case_set_id"])
    assert created.name == "新导入集" and created.revision == 1
    assert result["imported_count"] == result["generated_count"] == 1
    assert [row.name for row in cases_db.query(CaseItem).filter_by(case_set_id=created.id)] == ["登录成功"]
    assert {log.action for log in cases_db.query(AuditLog).filter_by(target_id=created.id)} == {
        "case_set_create", "case_set_import",
    }


def test_invalid_excel_does_not_create_empty_case_set(cases_db):
    """无效 Excel 不会提前创建空集；重复尝试也不会增加草稿。"""
    for _ in range(2):
        with pytest.raises(AppError) as error:
            asyncio.run(routes.create_case_set_from_excel(
                _request(), UploadFile(file=BytesIO(b"invalid"), filename="坏文件.xlsx"),
                "不应出现", None, cases_db, _actor(),
            ))
        assert error.value.code == ErrorCode.VALIDATION
    assert cases_db.query(CaseSet).count() == 1
    assert cases_db.query(AuditLog).count() == 0


def test_new_set_import_rolls_back_after_write_failure(cases_db, monkeypatch):
    """即使写入行时失败，新建集和相关审计也一起回滚。"""
    workbook = Workbook()
    workbook.active.append(["用例名称", "预期结果"])
    workbook.active.append(["登录成功", "进入首页"])
    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    write = routes._write_imported_cases

    def fail_after_rows(*args, **kwargs):
        write(*args, **kwargs)
        raise RuntimeError("数据库提交前失败")

    monkeypatch.setattr(routes, "_write_imported_cases", fail_after_rows)
    with pytest.raises(RuntimeError):
        asyncio.run(routes.create_case_set_from_excel(
            _request(), UploadFile(file=buffer, filename="需求.xlsx"),
            "不应出现", None, cases_db, _actor(),
        ))
    assert cases_db.query(CaseSet).count() == 1
    assert cases_db.query(CaseItem).count() == 0
    assert cases_db.query(AuditLog).count() == 0


def test_stale_replace_import_preserves_newer_cases(cases_db):
    """旧页面的整集替换导入不能清除另一个页面新保存的用例。"""
    routes.save_cases("set", CasesPayload(expected_revision=0, cases=[
        {"name": "新保存", "strategy": "正向", "priority": "HX", "expected": "成功"},
    ]), cases_db, _actor())
    workbook = Workbook()
    workbook.active.append(["用例名称", "预期结果"])
    workbook.active.append(["旧导入", "旧结果"])
    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    with pytest.raises(AppError) as error:
        asyncio.run(routes.import_case_set(
            "set", _request(), UploadFile(file=buffer, filename="旧页面.xlsx"),
            "replace", cases_db, _actor(), expected_revision=0,
        ))
    assert error.value.code == ErrorCode.CONCURRENCY
    assert [case.name for case in cases_db.query(CaseItem).all()] == ["新保存"]


@pytest.mark.parametrize("mode", ["append", "replace"])
@pytest.mark.parametrize("identity", ["new", "existing", "foreign"])
def test_import_repeated_ids_update_one_local_row(cases_db, mode, identity):
    """批内重复编号沿用同编号更新语义，跨集编号只重建一个本集副本。"""
    if identity == "existing":
        _add_case(cases_db, "repeated")
    elif identity == "foreign":
        cases_db.add(CaseSet(id="other-set", name="其他集合", created_by="other"))
        _add_case(cases_db, "repeated", case_set_id="other-set", name="其他集合原文")
    workbook = Workbook()
    workbook.active.append(["用例编号", "用例名称", "预期结果"])
    workbook.active.append(["repeated", "第一次", "首次预期"])
    workbook.active.append(["repeated", "最后一次", "最终预期"])
    buffer = BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    result = asyncio.run(routes.import_case_set(
        "set", _request(), UploadFile(file=buffer, filename="重复编号.xlsx"), mode, cases_db, _actor(),
        expected_revision=0,
    ))
    rows = cases_db.query(CaseItem).filter(CaseItem.case_set_id == "set").all()
    assert result["generated_count"] == len(rows) == 1
    assert rows[0].name == "最后一次" and rows[0].expected == "最终预期"
    if identity == "foreign":
        assert rows[0].id != "repeated"
        assert cases_db.get(CaseItem, "repeated").name == "其他集合原文"
