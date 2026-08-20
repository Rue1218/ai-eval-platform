"""Excel 用例导入解析与模板生成（PRD 5.4.1）。

支持三种表头，按列名别名识别，不要求列顺序与平台导出完全一致：

1. ``platform``：平台导出列（策略 / 优先级 / 模块 / 用例名称 / 前置条件 / 步骤 / 预期结果 / 测试类型）
2. ``standard``：testcase-tools 标准列（用例编号 / 所属模块 / 用例标题 / 优先级 / 用例类型 / 前置条件 / 测试步骤 / 预期结果）
3. ``simple``：简化列（模块 / 名称 / 步骤 / 预期）

纯函数，供路由与单测共用，禁止在此访问数据库。
"""

from __future__ import annotations

import io
import re
from typing import Any
from zipfile import BadZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet.worksheet import Worksheet

from .errors import AppError, ErrorCode

# 导入规模闸门：与数据集上传同量级，避免误把生成上限 80 套到存量用例库
IMPORT_MAX_ROWS = 2_000
IMPORT_MAX_BYTES = 10 * 1024 * 1024

# 平台导出/模板的固定列：字段名 -> 中文列名；导入时亦按别名识别
EXPORT_FIXED_COLUMNS: list[tuple[str, str]] = [
    ("id", "用例编号"),
    ("strategy", "策略"),
    ("priority", "优先级"),
    ("module", "模块"),
    ("name", "用例名称"),
    ("precondition", "前置条件"),
    ("steps", "步骤"),
    ("expected", "预期结果"),
    ("test_type", "测试类型"),
    ("mapped", "已映射"),
    ("pending_complete", "待补全"),
]

# 导入时跳过的运行态列，避免把导出快照里的映射标记写回用例正文
_SKIP_FIELDS = frozenset({"mapped", "pending_complete"})

# 列名别名（去空白、全角空格后精确匹配）-> 规范字段
_HEADER_ALIASES: dict[str, str] = {
    "用例编号": "id",
    "编号": "id",
    "id": "id",
    "caseid": "id",
    "case_id": "id",
    "策略": "strategy",
    "测试策略": "strategy",
    "用例策略": "strategy",
    "strategy": "strategy",
    "优先级": "priority",
    "级别": "priority",
    "priority": "priority",
    "模块": "module",
    "所属模块": "module",
    "功能模块": "module",
    "module": "module",
    "用例名称": "name",
    "用例标题": "name",
    "名称": "name",
    "标题": "name",
    "name": "name",
    "title": "name",
    "前置条件": "precondition",
    "前置": "precondition",
    "precondition": "precondition",
    "步骤": "steps",
    "测试步骤": "steps",
    "操作步骤": "steps",
    "steps": "steps",
    "预期结果": "expected",
    "预期": "expected",
    "期望结果": "expected",
    "expected": "expected",
    "测试类型": "test_type",
    "用例类型": "test_type",
    "类型": "test_type",
    "test_type": "test_type",
    "已映射": "mapped",
    "mapped": "mapped",
    "待补全": "pending_complete",
    "pending_complete": "pending_complete",
}

_STRATEGY_ALIASES: dict[str, str] = {
    "正向": "正向",
    "正向用例": "正向",
    "positive": "正向",
    "反向": "反向",
    "反向用例": "反向",
    "异常": "反向",
    "negative": "反向",
    "边界": "边界",
    "边界值": "边界",
    "boundary": "边界",
    "等价": "等价类",
    "等价类": "等价类",
    "equivalence": "等价类",
    "状态": "状态迁移",
    "状态迁移": "状态迁移",
    "state": "状态迁移",
    "场景": "场景",
    "场景用例": "场景",
    "scenario": "场景",
}

_PRIORITY_ALIASES: dict[str, str] = {
    "p0": "P0",
    "p1": "P1",
    "p2": "P2",
    "p3": "P3",
    "0": "P0",
    "1": "P1",
    "2": "P2",
    "3": "P3",
    "高": "P0",
    "中": "P1",
    "低": "P2",
}

_INSTRUCTION_SHEET = "填写说明"
_XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def xlsx_media_type() -> str:
    """Excel 媒体类型，导出与模板下载共用。"""
    return _XLSX_MEDIA_TYPE


def _norm_header(value: Any) -> str:
    """表头归一：转字符串、去空白与全角空格、小写英文字母保留中文。"""
    text = str(value or "").replace("\u3000", " ").strip()
    text = re.sub(r"\s+", "", text)
    return text.lower()


def _cell_text(value: Any) -> str:
    """单元格转可读字符串；布尔用是/否，其余去首尾空白。"""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "是" if value else "否"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _map_headers(cells: list[Any]) -> dict[int, str]:
    """把一行单元格映射为 {列下标: 规范字段}；无法识别的列不进入映射。"""
    mapping: dict[int, str] = {}
    for index, cell in enumerate(cells):
        key = _HEADER_ALIASES.get(_norm_header(cell))
        if key and key not in mapping.values():
            mapping[index] = key
    return mapping


def _detect_format(field_set: set[str]) -> str:
    """根据已识别字段判断导入格式标签，仅用于审计与前端提示。"""
    if "strategy" in field_set and "name" in field_set:
        return "platform"
    if "name" in field_set and "steps" in field_set and "priority" in field_set:
        return "standard"
    return "simple"


def _normalize_strategy(raw: str) -> str:
    """策略别名归一到六策略中文名；无法识别时默认正向。"""
    key = _norm_header(raw)
    return _STRATEGY_ALIASES.get(key) or _STRATEGY_ALIASES.get(raw.strip()) or "正向"


def _normalize_priority(raw: str) -> str:
    """优先级归一到 P0–P3；空值默认 P1。"""
    text = _norm_header(raw)
    if not text:
        return "P1"
    if text.upper() in {"P0", "P1", "P2", "P3"}:
        return text.upper()
    return _PRIORITY_ALIASES.get(text, "P1")


def _row_values(sheet: Worksheet, row_idx: int) -> list[Any]:
    """读取指定行全部单元格值（含空列，便于按下标对齐表头）。"""
    return [cell.value for cell in sheet[row_idx]]


def _find_header_row(sheet: Worksheet) -> tuple[int, dict[int, str]] | None:
    """在前 20 行内寻找可识别表头：至少命中名称列，且另外再命中一列业务字段。"""
    max_row = min(sheet.max_row or 0, 20)
    for row_idx in range(1, max_row + 1):
        mapping = _map_headers(_row_values(sheet, row_idx))
        fields = set(mapping.values()) - _SKIP_FIELDS
        if "name" in fields and len(fields) >= 2:
            return row_idx, mapping
    return None


def parse_cases_xlsx(content: bytes) -> tuple[list[dict[str, Any]], str, int]:
    """解析 xlsx 字节流为用例字典列表。

    返回 ``(cases, format, skipped_count)``。无表头、空文件、行数超限均抛 ``AppError(VALIDATION)``。
    扩展列（未识别表头）写入每条用例的 ``extras``。
    """
    try:
        workbook = load_workbook(io.BytesIO(content), data_only=True)
    except (InvalidFileException, BadZipFile, OSError, ValueError, KeyError) as exc:
        raise AppError(ErrorCode.VALIDATION, "无法解析 Excel，请使用 .xlsx 文件") from exc

    try:
        header_hit: tuple[Worksheet, int, dict[int, str]] | None = None
        for sheet in workbook.worksheets:
            if (sheet.title or "").strip() == _INSTRUCTION_SHEET:
                continue
            found = _find_header_row(sheet)
            if found:
                header_hit = (sheet, found[0], found[1])
                break
        if header_hit is None:
            raise AppError(
                ErrorCode.VALIDATION,
                "未识别到用例表头。请使用模板列，或提供「模块 / 名称 / 步骤 / 预期」简化列",
            )
        sheet, header_row, mapping = header_hit
        header_cells = _row_values(sheet, header_row)
        extra_cols = {
            index: _cell_text(header_cells[index]) or f"col_{index}"
            for index in range(len(header_cells))
            if index not in mapping and _cell_text(header_cells[index])
        }
        fmt = _detect_format(set(mapping.values()) - _SKIP_FIELDS)
        cases: list[dict[str, Any]] = []
        skipped = 0
        for row_idx in range(header_row + 1, (sheet.max_row or header_row) + 1):
            values = _row_values(sheet, row_idx)
            if not any(_cell_text(cell) for cell in values):
                continue
            item: dict[str, Any] = {}
            for index, field in mapping.items():
                if field in _SKIP_FIELDS:
                    continue
                raw = _cell_text(values[index] if index < len(values) else None)
                item[field] = raw
            extras: dict[str, Any] = {}
            for index, title in extra_cols.items():
                raw = _cell_text(values[index] if index < len(values) else None)
                if raw:
                    extras[title] = raw
            if extras:
                item["extras"] = extras
            name = str(item.get("name") or "").strip()
            if not name:
                skipped += 1
                continue
            item["name"] = name[:200]
            item["strategy"] = _normalize_strategy(str(item.get("strategy") or ""))
            item["priority"] = _normalize_priority(str(item.get("priority") or ""))
            item["module"] = str(item.get("module") or "")[:100]
            item["precondition"] = str(item.get("precondition") or "")
            item["steps"] = str(item.get("steps") or "")
            item["expected"] = str(item.get("expected") or "")
            item["test_type"] = str(item.get("test_type") or "")[:100]
            case_id = str(item.get("id") or "").strip()
            if case_id:
                item["id"] = case_id[:64]
            else:
                item.pop("id", None)
            cases.append(item)
            if len(cases) > IMPORT_MAX_ROWS:
                raise AppError(ErrorCode.VALIDATION, f"导入行数超过上限（{IMPORT_MAX_ROWS} 条）")
        if not cases:
            raise AppError(ErrorCode.VALIDATION, "Excel 中没有含用例名称的有效行")
        return cases, fmt, skipped
    finally:
        workbook.close()


def build_import_template() -> bytes:
    """生成带表头、示例行与填写说明的 Excel 模板。"""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "用例集"
    headers = [title for _field, title in EXPORT_FIXED_COLUMNS if _field not in _SKIP_FIELDS]
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="1F4E79")
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
    sheet.append(
        [
            "",
            "正向",
            "P0",
            "登录",
            "账密正确登录",
            "已注册账号处于正常状态",
            "1. 打开登录页\n2. 输入正确账密并提交",
            "进入工作台首页",
            "核心业务",
        ]
    )
    sheet.column_dimensions["A"].width = 16
    sheet.column_dimensions["B"].width = 12
    sheet.column_dimensions["C"].width = 10
    sheet.column_dimensions["D"].width = 14
    sheet.column_dimensions["E"].width = 28
    sheet.column_dimensions["F"].width = 28
    sheet.column_dimensions["G"].width = 36
    sheet.column_dimensions["H"].width = 28
    sheet.column_dimensions["I"].width = 14

    help_sheet = workbook.create_sheet(_INSTRUCTION_SHEET)
    help_sheet["A1"] = "用例 Excel 导入说明"
    help_sheet["A1"].font = Font(bold=True, size=14)
    help_lines = [
        "1. 请使用「用例集」工作表填写数据；本说明表导入时会被忽略。",
        "2. 平台标准列：用例编号、策略、优先级、模块、用例名称、前置条件、步骤、预期结果、测试类型。",
        "3. 亦支持 testcase-tools 标准列：用例编号、所属模块、用例标题、优先级、用例类型、前置条件、测试步骤、预期结果。",
        "4. 简化列只需：模块、名称、步骤、预期；缺策略时默认「正向」，缺优先级时默认 P1。",
        "5. 用例名称必填；空行跳过。表头可出现在前 20 行（允许首行写标题）。",
        "6. 策略取值：正向 / 反向 / 边界 / 等价类 / 状态迁移 / 场景；优先级 P0–P3。",
        "7. 未识别的额外列会写入用例扩展字段。已映射、待补全列为导出快照，导入时忽略。",
        "8. 请另存为 .xlsx；旧版 .xls 需先转换。单文件不超过 10MB、有效用例不超过 2000 条。",
    ]
    for index, line in enumerate(help_lines, start=3):
        help_sheet[f"A{index}"] = line
    help_sheet.column_dimensions["A"].width = 96

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
